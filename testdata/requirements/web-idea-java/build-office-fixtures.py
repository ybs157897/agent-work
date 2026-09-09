#!/usr/bin/env python3
"""Build the small DOCX requirement fixture used by parser acceptance tests."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


MARKER = "模拟测试材料｜未获用户确认"
DOC_ID = "SRC-DOCX-01"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_border(cell, color: str = "D9DEE5") -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_run_font(run, name: str = "PingFang SC", size: float = 10.5, color: str = "1F2933") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, space_after: float = 6, line_spacing: float = 1.15) -> None:
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = line_spacing


def add_text(doc, text: str, *, bold: bool = False, size: float = 10.5, color: str = "1F2933", after: float = 6):
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, after)
    run = paragraph.add_run(text)
    run.bold = bold
    set_run_font(run, size=size, color=color)
    return paragraph


def add_heading(doc, text: str, level: int = 1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    run.bold = True
    set_run_font(run, size=15 if level == 1 else 11.5, color="172B4D")
    return paragraph


def add_bullet(doc, text: str):
    paragraph = doc.add_paragraph(style="List Bullet")
    style_paragraph(paragraph, space_after=4)
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_fact_block(doc):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, "EEF3F8")
    set_cell_border(cell, "B8C7D9")
    paragraph = cell.paragraphs[0]
    style_paragraph(paragraph, space_after=0, line_spacing=1.1)
    lead = paragraph.add_run("模拟范围建议：")
    lead.bold = True
    set_run_font(lead, size=10.5, color="173A5E")
    rest = paragraph.add_run("首期仅文件阅读、定义和引用；写入/重命名应用需单独确认")
    set_run_font(rest, size=10.5, color="173A5E")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_header_footer(section):
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_paragraph(paragraph, space_after=0)
    run = paragraph.add_run(f"{MARKER}  ·  {DOC_ID}")
    set_run_font(run, size=8.5, color="64748B")

    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(paragraph, space_after=0)
    run = paragraph.add_run(f"{DOC_ID}  |  {MARKER}")
    set_run_font(run, size=8, color="64748B")


def build(output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    add_header_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = "PingFang SC"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.font.size = Pt(10.5)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("Java 阅读范围讨论稿")
    run.bold = True
    set_run_font(run, size=22, color="111827")

    subtitle = doc.add_paragraph()
    style_paragraph(subtitle, space_after=14)
    run = subtitle.add_run("需求输入模拟件 · 文档 ID SRC-DOCX-01")
    set_run_font(run, size=10, color="64748B")

    add_text(doc, "本文用于验证附件文字、段落和范围条件的解析结果。文中内容是模拟测试材料，未获用户确认。", after=4)
    add_text(doc, "这不是现行业务规则。", bold=True, color="7C2D12", after=10)
    add_fact_block(doc)

    add_heading(doc, "范围建议")
    add_bullet(doc, "首期支持在锁定的 Workspace Root 内阅读文件。")
    add_bullet(doc, "首期支持 Java 定义和引用查询，结果可为零条或多条。")
    add_bullet(doc, "文件写入、重命名和一键应用需要单独确认。")
    add_bullet(doc, "外部依赖目录按只读方式展示，不把依赖修改纳入本模拟范围。")

    add_heading(doc, "异常")
    add_bullet(doc, "索引未就绪：显示等待状态，并允许稍后重试。")
    add_bullet(doc, "JDK 缺失：保留文件阅读，明确提示 Java 语义服务不可用。")
    add_bullet(doc, "跳转零结果或多结果：显示可解释的结果状态，不自动猜选目标。")
    add_bullet(doc, "外部依赖只读：允许查看来源，不提供写入入口。")

    doc.add_page_break()
    add_heading(doc, "待确认事项")
    add_text(doc, "以下问题保持未决，仅用于测试解析器对条件、范围和确认状态的识别。", after=8)
    add_bullet(doc, "是否将文件阅读、定义和引用列为首期交付范围。")
    add_bullet(doc, "写入和重命名是否始终要求独立确认。")
    add_bullet(doc, "索引未就绪、JDK 缺失和零/多结果是否使用统一错误文案。")
    add_bullet(doc, "外部依赖是否只允许只读访问。")

    add_heading(doc, "边界说明")
    add_text(doc, "模拟范围建议不代表现行业务规则，不代表已经批准的产品行为，也不代表系统已经支持这些能力。", after=8)
    add_text(doc, "状态：pending。用户确认后，需求方应另行记录确认内容和后续变更。", bold=True, color="7C2D12", after=4)

    doc.save(output)


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    build(args.output)
